"""Phase 2.2: NVD CVSS severity enrichment (NVD CVE API 2.0; never the live API in tests).

Canonical metadata stays authoritative for applicability, versions and statuses; NVD only
supplies the CVSS rating persisted with each finding at analysis time.
"""

import io
import itertools
import json
import sqlite3
import urllib.parse
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from nvd_fixtures import (
    REPORT_CVES,
    FakeNvd,
    cve_obj,
    cvss,
    make_client,
    nvd_cache_db,
    response,
    v2,
)
from openpyxl import load_workbook
from phase2_fixtures import REAL_REPORT, ScriptedSSH, make_metadata
from test_analysis import AUTH_FAILURE, BAD, BAD_IP, GOOD, GOOD_IP, sync
from test_tags import save
from test_web import upload

from ec2patcher.app import create_app
from ec2patcher.database import _MIGRATIONS, Database
from ec2patcher.services import analysis_service, cve_resolver, nvd, security_metadata
from ec2patcher.services import excel_export as xl
from ec2patcher.services.analysis_service import AnalysisService, summarize

CVE = "CVE-2026-1000"


def select(**metrics):
    return nvd.result_from_cve(cve_obj(CVE, **metrics))


def picked(result):
    return (result.source, result.version, result.score, result.severity)


# --- CVSS metric selection -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("key", "version", "score", "raw", "severity"),
    [
        ("cvssMetricV40", "4.0", 8.7, "HIGH", "High"),
        ("cvssMetricV31", "3.1", 9.8, "CRITICAL", "Critical"),
        ("cvssMetricV30", "3.0", 5.3, "MEDIUM", "Medium"),
    ],
)
def test_single_nist_metric(key, version, score, raw, severity):
    r = select(**{key: [cvss(version, score, raw)]})
    assert r.status == nvd.OK
    assert picked(r) == ("nvd@nist.gov", version, score, severity)
    assert r.source_type == "Primary" and r.vector.startswith(f"CVSS:{version}/")
    assert r.last_modified == "2026-09-01T12:17:13.423" and r.note is None


def test_case_a_nvd_source_beats_newer_cna_version():
    r = select(
        cvssMetricV40=[cvss("4.0", 9.8, "CRITICAL", source="cna@vendor.example")],
        cvssMetricV31=[cvss("3.1", 8.8, "HIGH")],
    )
    assert picked(r) == ("nvd@nist.gov", "3.1", 8.8, "High")


def test_case_b_nvd_v40_beats_nvd_v31():
    r = select(
        cvssMetricV31=[cvss("3.1", 9.8, "CRITICAL")], cvssMetricV40=[cvss("4.0", 8.6, "HIGH")]
    )
    assert picked(r) == ("nvd@nist.gov", "4.0", 8.6, "High")


def test_case_c_primary_cna_v40_beats_secondary_v31():
    r = select(
        cvssMetricV40=[cvss("4.0", 6.9, "MEDIUM", source="cna@vendor.example")],
        cvssMetricV31=[cvss("3.1", 9.1, "CRITICAL", source="other@x.example", kind="Secondary")],
    )
    assert picked(r) == ("cna@vendor.example", "4.0", 6.9, "Medium")
    assert r.source_type == "Primary"


def test_case_d_primary_v31_beats_secondary_v40():
    r = select(
        cvssMetricV40=[cvss("4.0", 9.3, "CRITICAL", source="sec@x.example", kind="Secondary")],
        cvssMetricV31=[cvss("3.1", 7.5, "HIGH", source="cna@vendor.example")],
    )
    assert picked(r) == ("cna@vendor.example", "3.1", 7.5, "High")


def test_case_e_only_secondary_is_order_independent():
    metrics = {
        "cvssMetricV31": [
            cvss("3.1", 7.5, "HIGH", source="zeta@x.example", kind="Secondary"),
            cvss("3.1", 9.8, "CRITICAL", source="alpha@x.example", kind="Secondary"),
        ],
        "cvssMetricV30": [cvss("3.0", 5.0, "MEDIUM", source="aaa@x.example", kind="Secondary")],
        "cvssMetricV40": [
            cvss("4.0", 6.1, "MEDIUM", source="yak@x.example", kind="Secondary"),
            cvss("4.0", 8.2, "HIGH", source="bee@x.example", kind="Secondary"),
        ],
    }
    results = set()
    for order in itertools.permutations(metrics):
        for flip in (False, True):
            shuffled = {k: list(reversed(metrics[k])) if flip else metrics[k] for k in order}
            results.add(picked(select(**shuffled)))
    # Newest version first, then alphabetical source - never array order or highest score.
    assert results == {("bee@x.example", "4.0", 8.2, "High")}


def test_real_api_order_nvd_listed_second():
    """Real NVD responses often list the CNA metric first (e.g. CVE-2024-6387)."""
    r = select(
        cvssMetricV31=[
            cvss("3.1", 8.1, "HIGH", source="secalert@redhat.com", kind="Secondary"),
            cvss("3.1", 8.1, "HIGH"),
        ]
    )
    assert r.source == "nvd@nist.gov"


def test_nvd_source_is_case_insensitive_and_needs_no_type():
    r = select(
        cvssMetricV31=[
            cvss("3.1", 9.0, "CRITICAL", source="cna@x.example"),
            cvss("3.1", 4.0, "MEDIUM", source="NVD@NIST.GOV", kind=None),
        ]
    )
    assert (r.source, r.severity) == ("NVD@NIST.GOV", "Medium")


def test_missing_source_is_handled():
    data = {"version": "3.1", "baseScore": 7.1, "baseSeverity": "HIGH"}
    r = select(cvssMetricV31=[{"type": "Secondary", "cvssData": data}])
    assert (r.source, r.source_type, r.severity, r.vector) == (None, "Secondary", "High", None)


def test_missing_vector_kept_empty():
    r = select(cvssMetricV31=[cvss("3.1", 7.0, "HIGH", vector=False)])
    assert r.vector is None and r.severity == "High"


def test_missing_base_severity_derived_from_score():
    r = select(cvssMetricV31=[cvss("3.1", 9.1, None)])
    assert r.severity == "Critical" and "no baseSeverity" in r.note


def test_missing_score_is_never_invented():
    data = {"version": "3.1", "baseSeverity": "HIGH"}
    r = select(cvssMetricV31=[{"source": "nvd@nist.gov", "type": "Primary", "cvssData": data}])
    assert r.status == nvd.NO_CVSS and r.severity is None and r.score is None


@pytest.mark.parametrize(
    "metrics",
    [
        None,
        {},
        [],
        "garbage",
        {"cvssMetricV31": "not-a-list"},
        {"cvssMetricV31": [None, 5, "x", {"cvssData": None}, {"cvssData": []}]},
        {"cvssMetricV31": [{"cvssData": {"baseScore": "9.8"}}]},
        {"cvssMetricV31": [{"cvssData": {"baseScore": 11.0}}]},
        {"cvssMetricV31": [{"cvssData": {"baseScore": True}}]},
        {"ssvcV203": [{"source": "x"}]},
    ],
)
def test_malformed_or_missing_metrics_give_no_cvss(metrics):
    r = nvd.result_from_cve({"id": CVE, "metrics": metrics})
    assert r.status == nvd.NO_CVSS and r.severity is None and r.score is None


def test_malformed_metric_skipped_in_favour_of_valid_one():
    r = select(
        cvssMetricV40=[{"source": "nvd@nist.gov", "type": "Primary", "cvssData": {}}],
        cvssMetricV31=[cvss("3.1", 6.5, "MEDIUM")],
    )
    assert picked(r) == ("nvd@nist.gov", "3.1", 6.5, "Medium")


def test_inconsistent_base_severity_uses_score_and_records_it(caplog):
    r = select(cvssMetricV31=[cvss("3.1", 9.8, "LOW")])
    assert r.severity == "Critical"
    assert "'LOW' is inconsistent with base score 9.8" in r.note
    assert "inconsistent" in caplog.text


def test_v2_only_is_a_legacy_fallback_without_critical():
    r = select(cvssMetricV2=[v2(10.0, "HIGH")])
    assert picked(r) == ("nvd@nist.gov", "2.0", 10.0, "High") and r.note is None
    assert select(cvssMetricV2=[v2(5.0, "MEDIUM")]).severity == "Medium"


def test_v2_never_preferred_over_modern_metrics():
    r = select(
        cvssMetricV2=[v2(10.0, "HIGH")],
        cvssMetricV30=[cvss("3.0", 3.1, "LOW", source="cna@x.example", kind="Secondary")],
    )
    assert picked(r) == ("cna@x.example", "3.0", 3.1, "Low")


# --- severity normalization --------------------------------------------------------------


@pytest.mark.parametrize(
    ("score", "expected"),
    [
        (0.1, "Low"), (3.9, "Low"), (4.0, "Medium"), (6.9, "Medium"), (7.0, "High"),
        (8.9, "High"), (9.0, "Critical"), (10.0, "Critical"), (0.0, None),
    ],
)  # fmt: skip
@pytest.mark.parametrize("version", ["3.0", "3.1", "4.0"])
def test_severity_boundaries(score, expected, version):
    assert nvd.severity_for_score(score, version) == expected


def test_zero_score_is_not_low():
    r = select(cvssMetricV31=[cvss("3.1", 0.0, "NONE")])
    assert r.score == 0.0 and r.severity is None and "0.0 (None)" in r.note
    finding = cve_resolver.Finding(cve=CVE, source=None, status="NOT_AFFECTED", cvss=r)
    assert finding.severity == "Unknown"


# --- HTTP client -------------------------------------------------------------------------


def test_request_uses_cve_id_query_parameter_and_no_key(tmp_path):
    fake = FakeNvd({CVE: cve_obj(CVE, cvssMetricV31=[cvss("3.1", 7.5, "HIGH")])})
    client = make_client(tmp_path, fake)
    assert client.lookup(CVE.lower()).severity == "High"
    ((url, headers),) = fake.calls
    assert url == f"{nvd.API_URL}?cveId={CVE}"
    assert url.startswith("https://services.nvd.nist.gov/rest/json/cves/2.0?")
    assert "apiKey" not in headers and client.interval == nvd.INTERVAL_PUBLIC


def test_api_key_sent_as_header_only(tmp_path, monkeypatch):
    monkeypatch.setenv(nvd.API_KEY_ENV, "sekret-key")
    fake = FakeNvd({})
    client = make_client(tmp_path, fake)
    assert client.lookup(CVE).status == nvd.NOT_FOUND
    ((url, headers),) = fake.calls
    assert headers["apiKey"] == "sekret-key" and "sekret" not in url
    assert client.interval == nvd.INTERVAL_WITH_KEY
    with sqlite3.connect(nvd_cache_db(tmp_path)) as conn:
        assert "sekret" not in "\n".join(conn.iterdump())


def test_invalid_cve_id_is_not_sent(tmp_path):
    fake = FakeNvd({})
    r = make_client(tmp_path, fake).lookup("CVE-2026-1&apiKey=x")
    assert r.status == nvd.FAILED and fake.calls == []


def test_zero_vulnerabilities_is_not_found(tmp_path):
    r = make_client(tmp_path, FakeNvd({})).lookup(CVE)
    assert r.status == nvd.NOT_FOUND and r.severity is None


def test_invalid_json_fails_gracefully(tmp_path):
    r = make_client(tmp_path, FakeNvd(replies=[(200, {}, b"<html>oops")])).lookup(CVE)
    assert r.status == nvd.FAILED and "invalid JSON" in r.note


def test_unexpected_structure_fails_gracefully(tmp_path):
    r = make_client(tmp_path, FakeNvd(replies=[(200, {}, b'{"vulnerabilities": 3}')])).lookup(CVE)
    assert r.status == nvd.FAILED


def test_http_error_is_not_retried(tmp_path):
    fake = FakeNvd(replies=[(404, {}, b"")])
    r = make_client(tmp_path, fake).lookup(CVE)
    assert r.status == nvd.FAILED and r.note == "NVD HTTP 404" and len(fake.calls) == 1


def test_timeout_fails_gracefully_and_opens_circuit(tmp_path):
    fake = FakeNvd(replies=[TimeoutError("timed out")])
    client = make_client(tmp_path, fake)
    r = client.lookup(CVE)
    assert r.status == nvd.FAILED and "unreachable" in r.note
    # The next CVE of the same run does not wait for another timeout.
    assert client.lookup("CVE-2026-2000").status == nvd.FAILED
    assert len(fake.calls) == 1
    client.start_run()  # a new analysis run tries again
    fake.cves = {"CVE-2026-2000": cve_obj("CVE-2026-2000", cvssMetricV31=[cvss("3.1", 5, None)])}
    assert client.lookup("CVE-2026-2000").severity == "Medium"


def test_429_honours_retry_after(tmp_path):
    ok = (200, {}, response(cve_obj(CVE, cvssMetricV31=[cvss("3.1", 7.5, "HIGH")])))
    fake = FakeNvd(replies=[(429, {"Retry-After": "17"}, b""), ok])
    sleeps = []
    client = make_client(tmp_path, fake, sleep=sleeps.append, monotonic=lambda: 1000.0)
    assert client.lookup(CVE).severity == "High"
    assert 17.0 in sleeps and len(fake.calls) == 2


def test_5xx_retry_is_bounded_with_backoff(tmp_path):
    fake = FakeNvd(replies=[(503, {}, b"")] * 5)
    sleeps = []
    client = make_client(tmp_path, fake, sleep=sleeps.append, monotonic=lambda: 1000.0)
    r = client.lookup(CVE)
    assert r.status == nvd.FAILED and "503 after 3 attempts" in r.note
    assert len(fake.calls) == nvd.MAX_ATTEMPTS
    assert [s for s in sleeps if s >= 12] == [12.0, 24.0]  # 6 s interval x 2^attempt


def test_5xx_then_success(tmp_path):
    ok = (200, {}, response(cve_obj(CVE, cvssMetricV40=[cvss("4.0", 9.3, "CRITICAL")])))
    r = make_client(tmp_path, FakeNvd(replies=[(500, {}, b""), ok])).lookup(CVE)
    assert r.severity == "Critical" and r.version == "4.0"


def test_requests_are_spaced_by_rate_limit(tmp_path):
    clock = [100.0]
    sleeps = []

    def sleep(seconds):
        sleeps.append(seconds)
        clock[0] += seconds

    client = make_client(tmp_path, FakeNvd({}), sleep=sleep, monotonic=lambda: clock[0])
    for n in range(3):
        client.lookup(f"CVE-2026-{3000 + n}")
    assert sleeps == [6.0, 6.0]


# --- cache -------------------------------------------------------------------------------


class Clock:
    def __init__(self):
        self.value = datetime(2026, 9, 27, 12, 0, tzinfo=timezone.utc)

    def __call__(self):
        return self.value


def cached_client(tmp_path, fake, clock):
    return make_client(tmp_path, fake, now=clock)


def test_stale_cache_used_when_nvd_fails(tmp_path):
    clock = Clock()
    fake = FakeNvd({CVE: cve_obj(CVE, cvssMetricV31=[cvss("3.1", 7.5, "HIGH")])})
    cached_client(tmp_path, fake, clock).lookup(CVE)
    clock.value += timedelta(days=31)
    r = cached_client(tmp_path, FakeNvd(replies=[OSError("down")]), clock).lookup(CVE)
    assert (r.status, r.severity, r.score) == (nvd.STALE, "High", 7.5)
    assert "using cached data from 2026-09-27T12:00:00+00:00" in r.note


def test_no_cache_and_failure_is_unknown(tmp_path):
    r = make_client(tmp_path, FakeNvd(replies=[OSError("down")])).lookup(CVE)
    assert r.status == nvd.FAILED and r.severity is None
    finding = cve_resolver.Finding(cve=CVE, source="x", status="PATCH_AVAILABLE", cvss=r)
    assert finding.severity == "Unknown" and finding.status == "PATCH_AVAILABLE"


def test_not_found_is_cached_too(tmp_path):
    fake = FakeNvd({})
    make_client(tmp_path, fake).lookup(CVE)
    assert make_client(tmp_path, fake).lookup(CVE).status == nvd.NOT_FOUND
    assert len(fake.calls) == 1


def test_repeated_cve_in_one_run_is_looked_up_once_even_without_cache(tmp_path):
    fake = FakeNvd(replies=[(404, {}, b"")] * 5)
    client = make_client(tmp_path, fake)
    for _ in range(3):
        assert client.lookup(CVE).status == nvd.FAILED
    assert len(fake.calls) == 1


def test_unexpected_transport_bug_never_raises(tmp_path):
    def broken(url, headers, timeout):
        raise ZeroDivisionError("bug")

    r = make_client(tmp_path, broken).lookup(CVE)
    assert r.status == nvd.FAILED and "bug" in r.note


# --- analysis: capture, dedup, persistence --------------------------------------------------

BOTH = {GOOD: REAL_REPORT[GOOD], BAD: ["CVE-2026-63076", "CVE-2026-63075"]}


@pytest.fixture
def servers(db, pem_file):
    db.create_server(GOOD, GOOD_IP, str(pem_file))
    db.create_server(BAD, BAD_IP, str(pem_file))
    return db


def analyze(db, tmp_path, fake, report=BOTH, ssh=None, nvd_dir=None):
    db.save_report("security-report.json", report, "VALID")
    client = make_client(nvd_dir or tmp_path, fake)
    service = AnalysisService(
        db, make_metadata(tmp_path), runner=ssh or ScriptedSSH(), starter=sync,
        nvd_client=client,
    )  # fmt: skip
    run_id = service.start(db.get_latest_report())
    return db.get_analysis_run(run_id)


def test_same_cve_on_multiple_servers_is_one_nvd_call(servers, tmp_path):
    fake = FakeNvd(REPORT_CVES)
    run = analyze(servers, tmp_path, fake)
    assert [s.status for s in run.servers] == ["complete", "complete"]
    assert sorted(fake.requested()) == sorted(set(BOTH[GOOD]) | set(BOTH[BAD]))
    for s in run.servers:
        for f in s.findings:
            if f.cve == "CVE-2026-63076":
                assert (f.severity, f.cvss_score, f.cvss_source) == ("High", 7.5, "nvd@nist.gov")


def test_second_run_uses_cache(servers, tmp_path):
    fake = FakeNvd(REPORT_CVES)
    analyze(servers, tmp_path, fake)
    calls = len(fake.calls)
    run = analyze(servers, tmp_path, fake)
    assert len(fake.calls) == calls
    assert run.servers[0].findings[0].nvd_status == nvd.OK


def test_failed_server_makes_no_nvd_calls(servers, tmp_path):
    fake = FakeNvd(REPORT_CVES)
    run = analyze(servers, tmp_path, fake, report={BAD: ["CVE-2026-63076"]},
                  ssh=ScriptedSSH(failures=AUTH_FAILURE))  # fmt: skip
    assert run.servers[0].status == "failed" and fake.calls == []


def test_snapshot_persisted_with_priority(servers, tmp_path):
    run = analyze(servers, tmp_path, FakeNvd(REPORT_CVES), report={GOOD: REAL_REPORT[GOOD]})
    by_key = {(f.cve, f.source_package): f for f in run.servers[0].findings}
    kernel = by_key[("CVE-2026-54874", "linux-aws")]
    assert (kernel.cvss_severity, kernel.cvss_score, kernel.cvss_version) == (
        "Critical",
        9.3,
        "4.0",
    )
    assert (kernel.cvss_source, kernel.cvss_source_type) == ("security@kernel.org", "Primary")
    assert kernel.cvss_vector.startswith("CVSS:4.0/")
    assert kernel.nvd_last_modified == "2026-09-01T12:17:13.423" and kernel.nvd_status == "ok"
    assert kernel.cvss_label == "9.3 (v4.0)"
    # Canonical priority column untouched (the fixture metadata carries Medium for kernels).
    with sqlite3.connect(servers.path) as conn:
        cols = [r[1] for r in conn.execute("PRAGMA table_info(cve_findings)")]
    assert "priority" in cols and {"cvss_severity", "cvss_score", "nvd_status"} <= set(cols)
    # Reopening the database preserves everything.
    again = Database(servers.path).get_server_analysis(run.servers[0].id)
    assert again.findings == run.servers[0].findings


def test_nvd_failure_never_changes_patch_analysis(servers, tmp_path):
    ok = analyze(servers, tmp_path, FakeNvd(REPORT_CVES), nvd_dir=tmp_path / "a")
    down = analyze(
        servers, tmp_path, FakeNvd(replies=[OSError("no route")]), nvd_dir=tmp_path / "b"
    )

    def canonical(run):
        return [
            [(f.cve, f.source_package, f.installed_version, f.fixed_version, f.status, f.detail,
              f.binary_packages, f.pocket, f.priority) for f in s.findings]
            for s in run.servers
        ]  # fmt: skip

    def plans(run):
        return [[(p.binary_package, p.target_version, p.deb_filename, p.cves) for p in s.plan]
                for s in run.servers]  # fmt: skip

    assert canonical(ok) == canonical(down) and plans(ok) == plans(down)
    assert any(f.status == "PATCH_AVAILABLE" for f in down.servers[0].findings)
    assert {f.severity for s in down.servers for f in s.findings} == {"Unknown"}
    assert {f.nvd_status for s in down.servers for f in s.findings} == {nvd.FAILED}
    assert down.status == "completed"


def test_historical_snapshot_unchanged_when_nvd_changes(servers, tmp_path):
    run = analyze(servers, tmp_path, FakeNvd(REPORT_CVES))
    before = Database(servers.path).get_analysis_run(run.id).servers[0].findings
    changed = {c: cve_obj(c, cvssMetricV31=[cvss("3.1", 1.0, "LOW")]) for c in REPORT_CVES}
    new_run = analyze(servers, tmp_path, FakeNvd(changed), nvd_dir=tmp_path / "fresh-cache")
    assert Database(servers.path).get_analysis_run(run.id).servers[0].findings == before
    assert {f.severity for f in new_run.servers[0].findings} == {"Low"}


def test_progress_reports_nvd_stage(servers, tmp_path, monkeypatch):
    messages = []
    original = Database.update_analysis_run

    def spy(self, run_id, **fields):
        if fields.get("progress_message"):
            messages.append(fields["progress_message"])
        return original(self, run_id, **fields)

    monkeypatch.setattr(Database, "update_analysis_run", spy)
    analyze(servers, tmp_path, FakeNvd(REPORT_CVES))
    assert f"{GOOD} (1 of 2): Fetching NVD severity data (1 of 3)" in messages
    assert f"{BAD} (2 of 2): Fetching NVD severity data (2 of 2)" in messages


def test_severity_counters_independent_of_status(servers, tmp_path):
    run = analyze(servers, tmp_path, FakeNvd(REPORT_CVES), report={GOOD: [*REAL_REPORT[GOOD],
                  "CVE-2026-10004", "CVE-2026-10005"]})  # fmt: skip
    summary = summarize(run.servers[0])
    assert summary.by_severity == {"Critical": 1, "High": 1, "Medium": 1, "Low": 1, "Unknown": 1}
    assert sum(summary.by_severity.values()) == summary.reported == sum(summary.by_status.values())


def test_counters_for_failed_and_old_analyses_are_unknown(servers, tmp_path):
    run = analyze(servers, tmp_path, FakeNvd(REPORT_CVES), report={BAD: ["CVE-2026-63076"]},
                  ssh=ScriptedSSH(failures=AUTH_FAILURE))  # fmt: skip
    assert summarize(run.servers[0]).by_severity["Unknown"] == 1


# --- GUI ---------------------------------------------------------------------------------


@pytest.fixture
def web(db_path, tmp_path):
    def factory(nvd_transport=None, ssh=None):
        app = create_app(
            db_path=db_path, ssh_runner=ssh or ScriptedSSH(), metadata=make_metadata(tmp_path),
            analysis_starter=sync, shutdown_handler=lambda: None,
            nvd_client=make_client(tmp_path, nvd_transport or FakeNvd(REPORT_CVES)),
        )  # fmt: skip
        return TestClient(app, base_url="http://127.0.0.1")

    return factory


FULL = {GOOD: [*REAL_REPORT[GOOD], "CVE-2026-10004", "CVE-2026-10005"]}


def analyzed_page(web, pem_file, db_path, nvd_transport=None):
    with web(nvd_transport) as c:
        save(c, GOOD, GOOD_IP, pem_file)
        upload(c, FULL)
        c.post("/reports/analyze")
        run = Database(db_path).get_latest_analysis_run(details=True)
        return run, c.get(f"/analysis/{run.id}/servers/{run.servers[0].id}").text


def test_report_shows_nvd_severity_cvss_and_ubuntu_priority(web, pem_file, db_path):
    run, page = analyzed_page(web, pem_file, db_path)
    assert "<th>Severity</th><th>CVSS</th><th>Ubuntu Source Package</th>" in page
    for severity, css in [("Critical", "badge-critical"), ("High", "badge-danger"),
                          ("Medium", "badge-warning"), ("Low", "badge-low"),
                          ("Unknown", "badge-neutral")]:  # fmt: skip
        assert f'<span class="badge {css}">{severity}</span>' in page
    assert '<span title="CVSS:3.1/AV:N/AC:L">7.5 (v3.1)</span>' in page
    assert '<span title="CVSS:4.0/AV:N/AC:L/AT:N">9.3 (v4.0)</span>' in page
    assert "Source: NVD" in page and "Source: security@kernel.org (Primary)" in page
    assert "CVE not found in NVD" in page and '<span class="muted">—</span>' in page
    assert "Ubuntu Priority: Medium" in page  # Canonical priority still visible
    assert "Severity (NVD CVSS, per reported CVE):" in page
    for label, n in [("Critical", 1), ("High", 1), ("Medium", 1), ("Low", 1), ("Unknown", 1)]:
        assert f">{label}: {n}</span>" in page


def test_report_renders_when_nvd_failed(web, pem_file, db_path):
    run, page = analyzed_page(web, pem_file, db_path, FakeNvd(replies=[OSError("down")]))
    assert run.servers[0].status == "complete"
    assert "NVD lookup failed" in page and ">Unknown: 5</span>" in page
    assert "Patch available" in page  # Canonical status unaffected


def test_report_shows_stale_cache_marker(web, pem_file, db_path, tmp_path):
    old = REPORT_CVES["CVE-2026-63076"]
    Database(nvd_cache_db(tmp_path)).put_nvd_cache(
        "CVE-2026-63076", json.dumps(old["metrics"]), old["lastModified"],
        "2020-01-01T00:00:00+00:00",
    )  # fmt: skip
    _, page = analyzed_page(web, pem_file, db_path, FakeNvd(replies=[OSError("down")]))
    assert "Source: NVD &middot; stale cache" in page
    assert "using cached data from 2020-01-01T00:00:00+00:00" in page


# --- Excel -------------------------------------------------------------------------------

REQUIRED = [
    "CVE", "Severity", "CVSS Score", "CVSS Version", "Severity Source", "CVSS Vector",
    "Ubuntu Priority", "Ubuntu Source Package", "Installed Version", "Fixed Version", "Status",
    "Related Binary Package(s)",
]  # fmt: skip


def workbook_rows(content: bytes) -> tuple[list, list[dict]]:
    rows = list(load_workbook(io.BytesIO(content))["CVE Findings"].values)
    return list(rows[0]), [dict(zip(rows[0], r, strict=True)) for r in rows[1:]]


def test_excel_columns_match_persisted_snapshot(web, pem_file, db_path):
    run, _ = analyzed_page(web, pem_file, db_path)
    analysis = run.servers[0]
    with web() as c:
        r = c.get(f"/analysis/{run.id}/servers/{analysis.id}/export.xlsx")
    headers, rows = workbook_rows(r.content)
    assert all(h in headers for h in REQUIRED)
    assert headers.index("Severity") < headers.index("CVSS Score") < headers.index("Status")
    assert len(rows) == len(analysis.findings)
    for row, f in zip(rows, analysis.findings, strict=True):
        assert row["CVE"] == f.cve and row["Severity"] == f.severity
        assert row["CVSS Score"] == f.cvss_score
        assert (row["CVSS Version"] or None) == f.cvss_version
        assert (row["Severity Source"] or None) == f.cvss_source
        assert (row["CVSS Vector"] or None) == f.cvss_vector
        assert (row["Ubuntu Priority"] or None) == f.priority
    first = rows[0]
    assert (first["Severity"], first["CVSS Score"], first["CVSS Version"]) == ("High", 7.5, "3.1")
    assert first["Severity Source"] == "nvd@nist.gov" and first["NVD Lookup"] == "NVD"
    wb = load_workbook(io.BytesIO(r.content))
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(min_row=5, values_only=True)}
    assert summary["Severity Critical (NVD CVSS)"] == 1
    assert summary["Severity Unknown (NVD CVSS)"] == 1


def test_export_makes_zero_nvd_canonical_ssh_calls(web, pem_file, db_path, monkeypatch):
    run, _ = analyzed_page(web, pem_file, db_path)
    calls = []

    def forbidden(*args, **kwargs):
        calls.append(args)
        raise AssertionError("must not be called during export")

    for target, name in [
        (nvd.NvdClient, "lookup"), (nvd.NvdClient, "fetch"), (nvd, "http_get"),
        (security_metadata.SecurityMetadata, "lookup"),
        (security_metadata.SecurityMetadata, "ensure_fresh"),
        (analysis_service.ssh_service, "run_remote"),
        (analysis_service.AnalysisService, "run"),
    ]:  # fmt: skip
        monkeypatch.setattr(target, name, forbidden)

    def dump():
        with sqlite3.connect(db_path) as conn:
            return list(conn.iterdump())

    before = dump()
    fake, ssh = FakeNvd(REPORT_CVES), ScriptedSSH()
    with web(fake, ssh=ssh) as c:
        r = c.get(f"/analysis/{run.id}/servers/{run.servers[0].id}/export.xlsx")
    assert r.status_code == 200
    assert calls == [] and fake.calls == [] and ssh.calls == []
    assert dump() == before


def test_old_v3_findings_export_with_unknown_severity(db_path):
    db_path.parent.mkdir(parents=True)
    conn = sqlite3.connect(db_path)
    for version in (1, 2, 3):
        conn.executescript(_MIGRATIONS[version])
    conn.execute("PRAGMA user_version = 3")
    conn.execute(
        "INSERT INTO analysis_runs (report_filename, report_uploaded_at, report_content, "
        "started_at, status) VALUES ('old.json', 't', '{\"s\": [\"CVE-2025-2\"]}', "
        "'2026-01-02T03:04:05+00:00', 'completed')"
    )
    conn.execute(
        "INSERT INTO server_analyses (run_id, position, server_name, reported_cves, status) "
        "VALUES (1, 0, 's', '[\"CVE-2025-2\"]', 'complete')"
    )
    conn.execute(
        "INSERT INTO cve_findings (server_analysis_id, cve, source_package, status, priority) "
        "VALUES (1, 'CVE-2025-2', 'bash', 'ALREADY_FIXED', 'Negligible')"
    )
    conn.commit()
    conn.close()
    db = Database(db_path)
    run = db.get_analysis_run(1)
    (f,) = run.servers[0].findings
    assert (f.severity, f.cvss_label, f.priority, f.nvd_status) == (
        "Unknown",
        "—",
        "Negligible",
        None,
    )
    _, rows = workbook_rows(xl.build_workbook(run, run.servers[0]))
    (row,) = rows
    assert (row["Severity"], row["CVSS Score"], row["Ubuntu Priority"]) == (
        "Unknown",
        None,
        "Negligible",
    )
    assert row["NVD Lookup"].startswith("Not captured")


def test_url_encoding_is_parameterised():
    query = urllib.parse.urlsplit(f"{nvd.API_URL}?{urllib.parse.urlencode({'cveId': CVE})}").query
    assert urllib.parse.parse_qs(query) == {"cveId": [CVE]}
