"""Excel (.xlsx) export of one server's stored pre-patch report.

The workbook is built in memory from the persisted analysis snapshot only: it never re-runs
the analysis, opens ssh connections, reads Canonical metadata, queries NVD or talks to APT.
Values mirror the server report page so the GUI and the workbook always agree.
"""

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ec2patcher.formatting import format_size, format_timestamp
from ec2patcher.models import AnalysisRun, ServerAnalysis
from ec2patcher.services import analysis_service, cve_resolver, nvd
from ec2patcher.services.severity import SEVERITIES, UNKNOWN

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
NOT_ANALYZED_LABEL = "Not analyzed"
_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._-]+")

BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

# (header, column width); long text columns wrap.
CVE_COLUMNS = [
    ("CVE", 18),
    ("Severity", 11),
    ("CVSS Score", 11),
    ("CVSS Version", 12),
    ("Severity Source", 22),
    ("CVSS Vector", 44),
    ("Ubuntu Priority", 15),
    ("Ubuntu Source Package", 24),
    ("Installed Version", 26),
    ("Canonical Status", 22),
    ("Fixed Version", 26),
    ("APT Candidate", 40),
    ("Fix Pocket", 18),
    ("Status", 28),
    ("Related Binary Package(s)", 40),
    ("NVD Lookup", 30),
    ("Detail", 60),
]
PLAN_COLUMNS = [
    ("Binary Package", 30),
    ("Architecture", 12),
    ("Source Package", 22),
    ("Current Version", 24),
    ("Target Version", 24),
    ("Dependency", 11),
    (".deb Filename", 50),
    ("Download URI", 80),
    ("Size", 11),
    ("Size (bytes)", 13),
    ("Checksum", 50),
    ("Related CVEs", 22),
    ("Reboot Impact", 30),
    ("Plan Status", 12),
    ("Note", 50),
]


def analysis_time(run: AnalysisRun, analysis: ServerAnalysis) -> str:
    return analysis.completed_at or analysis.started_at or run.started_at


def export_filename(run: AnalysisRun, analysis: ServerAnalysis) -> str:
    """ec2patcher_<canonical server name>_<analysis date>.xlsx, safe for any filesystem."""
    name = _UNSAFE_FILENAME_CHARS.sub("_", analysis.server_name).strip("._")[:100] or "server"
    stamp = analysis_time(run, analysis)
    try:
        date = datetime.fromisoformat(stamp).astimezone().strftime("%Y-%m-%d")
    except (TypeError, ValueError):
        date = "unknown-date"
    return f"ec2patcher_{name}_{date}.xlsx"


def build_workbook(run: AnalysisRun, analysis: ServerAnalysis) -> bytes:
    """Return the .xlsx bytes for one stored server analysis (findings + plan loaded)."""
    wb = Workbook()
    _summary_sheet(wb.active, run, analysis)
    _table_sheet(wb.create_sheet("CVE Findings"), CVE_COLUMNS, _cve_rows(analysis))
    _table_sheet(wb.create_sheet("Package Plan"), PLAN_COLUMNS, _plan_rows(analysis))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- cells ---------------------------------------------------------------------------


def _set(ws: Worksheet, row: int, column: int, value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=column, value=value)
    if cell.data_type == "f":  # never let stored text become a formula
        cell.data_type = "s"
    return cell


def _yes_no_unknown(value: bool | None, yes: str, no: str) -> str:
    if value is None:
        return "Unknown"
    return yes if value else no


# --- Summary ---------------------------------------------------------------------------


def _summary_rows(run: AnalysisRun, analysis: ServerAnalysis) -> list[tuple[str, object]]:
    summary = analysis_service.summarize(analysis)
    rows: list[tuple[str, object]] = [("Server Name", analysis.server_name)]
    if analysis.display_name:
        rows.append(("Display Name", analysis.display_name))
    rows += [
        ("IP", analysis.ip_address or "-"),
        ("Remote Hostname", analysis.remote_hostname or "-"),
        ("Ubuntu", analysis.os_pretty_name or "-"),
        ("Ubuntu Version", analysis.os_version_id or "-"),
        ("Codename", analysis.os_codename or "-"),
        ("Architecture", analysis.architecture or "-"),
        ("Running Kernel", analysis.running_kernel or "-"),
        ("Analysis Status", analysis.status),
    ]
    if analysis.error:
        rows.append(("Analysis Error", analysis.error))
    rows += [
        ("Analysis Time", format_timestamp(analysis_time(run, analysis))),
        ("Analysis Run", f"#{run.id}"),
        ("Report File", run.report_filename),
    ]

    tally = run.lookup_tally
    if run.metadata_updated_at:
        metadata = "Online per-CVE lookup"
        if tally.total:
            metadata = f"{tally.state} per-CVE lookup · {tally.label}"
        if run.metadata_checked_at:
            metadata += f" · checked {format_timestamp(run.metadata_checked_at)}"
    else:
        metadata = "not available"
    rows += [
        ("Canonical Security Metadata", metadata),
        ("Canonical Metadata Source", run.metadata_source or "-"),
        ("Canonical Metadata Stale", "YES" if run.metadata_stale else "NO"),
    ]
    if run.metadata_warning:
        rows.append(("Canonical Metadata Warning", run.metadata_warning))

    if analysis.apt_updated_at:
        apt = f"Workstation private lists updated {format_timestamp(analysis.apt_updated_at)}"
        if analysis.apt_age_hours is not None:
            apt += f" ({analysis.apt_age_hours:.1f} hours before analysis)"
    else:
        apt = "not used"
    current_reboot = _yes_no_unknown(analysis.current_reboot_required, "YES", "NO")
    if analysis.current_reboot_required and analysis.reboot_required_packages:
        current_reboot += f" ({', '.join(analysis.reboot_required_packages)})"
    rows += [
        ("APT Metadata", apt),
        ("Current Reboot Required", current_reboot),
        (
            "Expected Reboot After Planned Patch",
            _yes_no_unknown(analysis.expected_reboot, "YES EXPECTED", "NO EXPECTED REBOOT"),
        ),
    ]
    if analysis.expected_reboot_reason:
        rows.append(("Expected Reboot Reason", analysis.expected_reboot_reason))
    rows += [
        ("Reboot Note", cve_resolver.REBOOT_HELP),
        ("Reported CVEs", summary.reported),
        *((f"Severity {s} (NVD CVSS)", summary.by_severity[s]) for s in SEVERITIES),
        *((f"{title} (CVEs)", n) for _, title, n in summary.buckets),
        ("Patch available", summary.count(cve_resolver.PATCH_AVAILABLE)),
        ("Already fixed", summary.count(cve_resolver.ALREADY_FIXED)),
        ("Not affected", summary.count(cve_resolver.NOT_AFFECTED)),
        ("Package not installed", summary.count(cve_resolver.PACKAGE_NOT_INSTALLED)),
        (
            "Published fix requires repository access",
            summary.count(
                cve_resolver.PRO_OR_ESM_REQUIRED,
                cve_resolver.FIX_NOT_IN_CONFIGURED_REPOS,
            ),
        ),
        (
            "No published fix / deferred",
            summary.count(cve_resolver.NO_FIX_PUBLISHED, cve_resolver.PENDING_OR_DEFERRED),
        ),
        (
            "Needs investigation / errors",
            summary.count(
                cve_resolver.UNKNOWN,
                cve_resolver.METADATA_UNAVAILABLE,
                cve_resolver.ANALYSIS_ERROR,
                "NOT_ANALYZED",
            ),
        ),
        ("Binary packages to update", summary.packages),
        (".deb files required", summary.debs),
        ("Unresolved plan entries", summary.unresolved),
        ("Estimated download size", format_size(summary.download_bytes)),
        ("Estimated download size (bytes)", summary.download_bytes),
    ]
    if analysis.warnings:
        rows.append(("Warnings", "\n".join(analysis.warnings)))
    return rows


def _summary_sheet(ws: Worksheet, run: AnalysisRun, analysis: ServerAnalysis) -> None:
    ws.title = "Summary"
    _set(ws, 1, 1, "EC2Patcher Pre-Patch Report").font = Font(bold=True, size=14)
    _set(
        ws, 2, 1, "Read-only snapshot of the stored analysis. Nothing was downloaded or installed."
    )
    _set(ws, 4, 1, "Field").font = BOLD
    _set(ws, 4, 2, "Value").font = BOLD
    for offset, (label, value) in enumerate(_summary_rows(run, analysis)):
        row = 5 + offset
        _set(ws, row, 1, label).font = BOLD
        _set(ws, row, 2, value).alignment = Alignment(
            wrap_text=True, vertical="top", horizontal="left"
        )
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 90


# --- tables ----------------------------------------------------------------------------


def _cve_rows(analysis: ServerAnalysis) -> list[list]:
    """Every stored finding (none collapsed), in report order; unanalyzed CVEs included."""
    by_cve: dict[str, list] = {}
    for f in analysis.findings:
        by_cve.setdefault(f.cve, []).append(f)
    order = [*analysis.reported_cves, *(c for c in by_cve if c not in analysis.reported_cves)]
    rows = []
    for cve in dict.fromkeys(order):
        findings = by_cve.get(cve)
        if not findings:
            rows.append([cve, UNKNOWN, *[""] * 11, NOT_ANALYZED_LABEL, "", "", ""])
            continue
        for f in findings:
            lookup = nvd.STATUS_LABELS.get(f.nvd_status, f.nvd_status) if f.nvd_status else ""
            if f.nvd_note:
                lookup = f"{lookup}: {f.nvd_note}" if lookup else f.nvd_note
            rows.append(
                [
                    f.cve,
                    f.severity,
                    f.cvss_score,
                    f.cvss_version or "",
                    f.cvss_source or "",
                    f.cvss_vector or "",
                    f.priority or "",
                    f.source_package or "",
                    f.installed_version or "",
                    f.canonical_status or "",
                    f.fixed_version or "",
                    f.apt_candidate or "",
                    f.pocket or "",
                    cve_resolver.STATUS_LABELS.get(f.status, f.status),
                    ", ".join(f.binary_packages),
                    lookup or "Not captured (analysis predates NVD severity)",
                    f.detail or "",
                ]
            )
    return rows


def _plan_rows(analysis: ServerAnalysis) -> list[list]:
    return [
        [
            p.binary_package,
            p.architecture,
            p.source_package or "",
            p.current_version or "not installed",
            p.target_version,
            "Yes" if p.is_dependency else "No",
            p.deb_filename or "",
            p.uri or "",
            format_size(p.size) if p.size is not None else "",
            p.size,
            p.checksum or "",
            ", ".join(p.cves) if p.cves else "required dependency",
            p.reboot_impact or "No reboot expected",
            p.status,
            p.reason or "",
        ]
        for p in analysis.plan
    ]


def _table_sheet(ws: Worksheet, columns: list[tuple[str, int]], rows: list[list]) -> None:
    for index, (header, width) in enumerate(columns, start=1):
        _set(ws, 1, index, header).font = BOLD
        ws.column_dimensions[get_column_letter(index)].width = width
    for r, values in enumerate(rows, start=2):
        for c, value in enumerate(values, start=1):
            _set(ws, r, c, value).alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"
